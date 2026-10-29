"""Carga inicial del catálogo desde Excel. Guía para el usuario: docs/carga-inicial.md."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Alias, Lote, Movimiento, Producto
from .texto import normalizar, normalizar_codigo

COLUMNAS = [
    # (nombre, obligatoria, descripción para la hoja de instrucciones)
    ("nombre", True, "Nombre tal como aparece impreso en el producto."),
    ("referencia", True, "Código interno o del proveedor. Único por producto."),
    ("codigo_barras", False, "Código de barras de fábrica, si lo tiene. Escríbalo como texto."),
    ("metodo", False, "Cómo se identifica: barras, nombre o qr. Si se deja vacío se deduce."),
    ("alias", False, "Otros nombres o abreviaturas, separados por coma."),
    ("presentacion", False, "Tamaño, concentración o unidad (ej. 500 mg x 30 tab)."),
    ("stock_minimo", True, "Cantidad mínima. Por debajo de ella se sugiere pedir."),
    ("cantidad_actual", True, "Inventario de arranque (número entero)."),
    ("fecha_vencimiento", False, "Vencimiento del inventario de arranque (AAAA-MM-DD)."),
    ("lote", False, "Lote del inventario de arranque."),
]
METODOS = {"barras", "nombre", "qr"}


def generar_plantilla() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Catalogo"
    verde = PatternFill("solid", fgColor="7ED957")
    gris = PatternFill("solid", fgColor="E5ECF1")
    for i, (col, obligatoria, _) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=col)
        celda.font = Font(bold=True, color="051923")
        celda.fill = verde if obligatoria else gris
        hoja.column_dimensions[celda.column_letter].width = max(16, len(col) + 4)
    hoja.append(["Acetaminofén", "MED-0001", "7702001000011", "barras", "Acetaminofen, Paracetamol",
                 "500 mg x 100 tab", 20, 35, "2027-06-30", "L2406A"])
    hoja.append(["Jeringa desechable", "INS-0002", "", "qr", "", "5 ml", 50, 120, "", ""])
    for col in ("B", "C"):  # referencias y códigos como texto para que Excel no los altere
        for fila in range(2, 2001):
            hoja[f"{col}{fila}"].number_format = "@"
    dv = DataValidation(type="list", formula1='"barras,nombre,qr"', allow_blank=True)
    hoja.add_data_validation(dv)
    dv.add("D2:D2000")
    hoja.freeze_panes = "A2"

    ayuda = wb.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Obligatoria", "Qué va"])
    for col, obligatoria, desc in COLUMNAS:
        ayuda.append([col, "Sí" if obligatoria else "No", desc])
    ayuda.append([])
    ayuda.append(["Las filas 2 y 3 de la hoja Catalogo son ejemplos: bórrelas antes de cargar."])
    for c in ("A", "B", "C"):
        ayuda.column_dimensions[c].width = 22 if c != "C" else 70
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _entero(valor) -> int | None:
    if valor is None or str(valor).strip() == "":
        return None
    try:
        numero = float(str(valor).replace(",", "."))
    except ValueError:
        return None
    return int(numero) if numero >= 0 and numero == int(numero) else None


def _fecha(valor) -> date | None | str:
    """Devuelve la fecha, None si está vacía o 'error' si no se entiende."""
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    return "error"


def leer_excel(contenido: bytes) -> tuple[list[dict], list[str]]:
    """Lee y valida el archivo. Devuelve (filas válidas, errores)."""
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        return [], ["El archivo no es un Excel válido (.xlsx)."]
    hoja = wb["Catalogo"] if "Catalogo" in wb.sheetnames else wb.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    try:
        encabezado = [normalizar(c).replace(" ", "_") if c else "" for c in next(filas)]
    except StopIteration:
        return [], ["El archivo está vacío."]

    faltan = [c for c, oblig, _ in COLUMNAS if oblig and c not in encabezado]
    if faltan:
        return [], [f"Faltan columnas obligatorias: {', '.join(faltan)}. Use la plantilla oficial."]
    pos = {c: encabezado.index(c) for c, _, _ in COLUMNAS if c in encabezado}

    validas, errores = [], []
    refs, codigos = {}, {}
    for n, fila in enumerate(filas, start=2):
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            continue

        def v(col):
            i = pos.get(col)
            return fila[i] if i is not None and i < len(fila) else None

        problemas = []
        nombre = str(v("nombre") or "").strip()
        ref = normalizar_codigo(v("referencia"))
        codigo = normalizar_codigo(v("codigo_barras")) or None
        metodo = str(v("metodo") or "").strip().lower()
        stock_min, cantidad = _entero(v("stock_minimo")), _entero(v("cantidad_actual"))
        vence = _fecha(v("fecha_vencimiento"))

        if not nombre:
            problemas.append("nombre vacío")
        if not ref:
            problemas.append("referencia vacía")
        elif ref in refs:
            problemas.append(f"referencia repetida (también en la fila {refs[ref]})")
        if codigo and codigo in codigos:
            problemas.append(f"código de barras repetido (también en la fila {codigos[codigo]})")
        if stock_min is None:
            problemas.append("stock_minimo debe ser un número entero mayor o igual a 0")
        if cantidad is None:
            problemas.append("cantidad_actual debe ser un número entero mayor o igual a 0")
        if metodo and metodo not in METODOS:
            problemas.append("metodo debe ser barras, nombre o qr")
        if metodo == "barras" and not codigo:
            problemas.append("metodo es barras pero no hay código de barras")
        if vence == "error":
            problemas.append("fecha_vencimiento no se entiende (use AAAA-MM-DD)")

        if problemas:
            errores.append(f"Fila {n}: " + "; ".join(problemas) + ".")
            continue
        refs[ref] = n
        if codigo:
            codigos[codigo] = n
        alias = [a.strip() for a in str(v("alias") or "").split(",") if a.strip()]
        validas.append({
            "fila": n, "nombre": nombre, "referencia": ref, "codigo_barras": codigo,
            "metodo": metodo or ("barras" if codigo else "nombre"), "alias": alias,
            "presentacion": str(v("presentacion") or "").strip(),
            "stock_minimo": stock_min, "cantidad_actual": cantidad,
            "fecha_vencimiento": vence, "lote": str(v("lote") or "").strip(),
        })
    return validas, errores


def aplicar_carga(db: Session, filas: list[dict], usuario: str) -> dict:
    """Crea o actualiza productos por referencia. Todo en una sola transacción."""
    creados = actualizados = 0
    # Conflictos de código de barras contra productos ya existentes con otra referencia.
    conflictos = []
    for f in filas:
        if f["codigo_barras"]:
            otro = db.scalar(select(Producto).where(Producto.codigo_barras == f["codigo_barras"]))
            if otro and otro.referencia != f["referencia"]:
                conflictos.append(f"Fila {f['fila']}: el código de barras ya pertenece a {otro.referencia}.")
    if conflictos:
        return {"creados": 0, "actualizados": 0, "errores": conflictos}

    for f in filas:
        p = db.scalar(select(Producto).where(Producto.referencia == f["referencia"]))
        nuevo = p is None
        if nuevo:
            p = Producto(referencia=f["referencia"], cantidad_actual=0)
            db.add(p)
        p.nombre = f["nombre"]
        p.nombre_normalizado = normalizar(f["nombre"])
        p.codigo_barras = f["codigo_barras"]
        p.metodo_identificacion = f["metodo"]
        p.presentacion = f["presentacion"]
        p.stock_minimo = f["stock_minimo"]
        db.flush()
        if not p.codigo_qr:
            p.codigo_qr = f"SNP-{p.id:06d}"

        existentes = {a.texto_normalizado for a in p.alias}
        for texto in f["alias"]:
            tn = normalizar(texto)
            if tn and tn not in existentes:
                p.alias.append(Alias(texto=texto, texto_normalizado=tn, origen="catalogo"))
                existentes.add(tn)

        if nuevo:
            creados += 1
            p.cantidad_actual = f["cantidad_actual"]
            db.add(Movimiento(producto_id=p.id, tipo="carga_inicial", cantidad_antes=0,
                              cantidad_despues=f["cantidad_actual"], detalle="Carga desde Excel",
                              usuario=usuario))
            if f["cantidad_actual"] > 0:
                db.add(Lote(producto_id=p.id, fecha_ingreso=date.today(),
                            fecha_vencimiento=f["fecha_vencimiento"], lote=f["lote"],
                            cantidad=f["cantidad_actual"], origen="carga_inicial", usuario=usuario))
        else:
            actualizados += 1  # en una recarga no se pisa el inventario ya registrado
    db.commit()
    return {"creados": creados, "actualizados": actualizados, "errores": []}
