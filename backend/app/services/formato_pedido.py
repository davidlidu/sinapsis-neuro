"""Formato de pedido en blanco (PDF con campos para llenar, o Excel) y su lectura de vuelta.

Sirve para quien prefiere armar el pedido a mano: lo imprime o lo llena en el computador y lo sube;
Sinapsis crea el borrador con esas cantidades. Las cantidades van en unidades (kits, cajas).
"""
import io
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Producto
from .texto import normalizar, normalizar_codigo

CAMPO = "pedir_"  # nombre de cada campo del PDF: pedir_<id del producto>
AZUL, VERDE, GRIS, TENUE = HexColor("#051923"), HexColor("#7ED957"), HexColor("#64748B"), HexColor("#E5ECF1")


def productos_del_formato(db: Session) -> list[Producto]:
    return list(db.scalars(select(Producto).where(Producto.activo.is_(True))
                           .order_by(Producto.seccion, Producto.nombre)))


def _minimo(p: Producto) -> str:
    if not p.stock_minimo:
        return ""
    return f"{p.stock_minimo:,} pruebas".replace(",", ".") if p.pruebas_por_unidad else f"{p.stock_minimo:,}".replace(",", ".")


def generar_pdf(productos: list[Producto]) -> bytes:
    salida = io.BytesIO()
    c = canvas.Canvas(salida, pagesize=letter, pageCompression=1)
    c.setTitle("Formato de pedido a farmacia · Sinapsis")
    ancho, alto = letter
    m = 36
    cols = [("Cód. farmacia", 72), ("Producto", 262), ("Presentación", 62), ("Mín.", 66), ("Cantidad a pedir", 78)]
    xs = [m]
    for _, w in cols:
        xs.append(xs[-1] + w)
    pagina = 0

    def encabezado():
        nonlocal pagina
        pagina += 1
        y = alto - m
        c.setFillColor(AZUL); c.setFont("Helvetica-Bold", 14)
        c.drawString(m, y - 4, "Formato de pedido a farmacia")
        c.setFont("Helvetica", 8); c.setFillColor(GRIS)
        c.drawRightString(ancho - m, y - 4, f"Sinapsis · {date.today().strftime('%d/%m/%Y')} · Página {pagina}")
        y -= 22
        if pagina == 1:
            c.setFillColor(AZUL); c.setFont("Helvetica", 9)
            c.drawString(m, y, "Solicitado por:")
            c.acroForm.textfield(name="solicitado_por", x=m + 70, y=y - 5, width=210, height=16, borderWidth=0.5,
                                 borderColor=GRIS, fillColor=HexColor("#FFFFFF"), fontSize=9)
            c.drawString(m + 300, y, "Fecha:")
            c.acroForm.textfield(name="fecha", x=m + 334, y=y - 5, width=100, height=16, borderWidth=0.5,
                                 borderColor=GRIS, fillColor=HexColor("#FFFFFF"), fontSize=9)
            y -= 20
            c.setFont("Helvetica", 7.5); c.setFillColor(GRIS)
            c.drawString(m, y, "Escriba la cantidad a pedir en unidades (kits, cajas, frascos). Deje vacío lo que no necesita.")
            c.drawString(m, y - 10, "Puede llenarlo en el computador y subirlo en Sinapsis → Pedidos → Cargar formato lleno, o imprimirlo.")
            y -= 24
        c.setFillColor(AZUL); c.rect(m, y - 14, xs[-1] - m, 16, fill=1, stroke=0)
        c.setFillColor(HexColor("#FFFFFF")); c.setFont("Helvetica-Bold", 7.5)
        for (titulo, _), x in zip(cols, xs):
            c.drawString(x + 4, y - 9, titulo)
        return y - 14

    y = encabezado()
    seccion = None
    for p in productos:
        nombre = simpleSplit(p.nombre, "Helvetica", 7.5, cols[1][1] - 8)[:2]
        alto_fila = 12 + 9 * len(nombre)
        nueva_seccion = (p.seccion or "Sin sección") != seccion
        if y - alto_fila - (16 if nueva_seccion else 0) < m + 10:
            c.showPage()
            y = encabezado()
            nueva_seccion = True
        if nueva_seccion:
            seccion = p.seccion or "Sin sección"
            c.setFillColor(TENUE); c.rect(m, y - 15, xs[-1] - m, 15, fill=1, stroke=0)
            c.setFillColor(AZUL); c.setFont("Helvetica-Bold", 8)
            c.drawString(m + 4, y - 11, seccion)
            y -= 15
        c.setFillColor(AZUL)
        c.setFont("Helvetica-Bold", 8.5); c.drawString(xs[0] + 4, y - 12, p.referencia)
        c.setFont("Helvetica", 7.5)
        for i, linea in enumerate(nombre):
            c.drawString(xs[1] + 4, y - 12 - 9 * i, linea)
        c.drawString(xs[2] + 4, y - 12, (p.presentacion or "")[:14])
        c.drawString(xs[3] + 4, y - 12, _minimo(p))
        c.acroForm.textfield(name=f"{CAMPO}{p.id}", x=xs[4] + 6, y=y - alto_fila + 4, width=cols[4][1] - 12,
                             height=min(alto_fila - 6, 16), borderWidth=0.5, borderColor=GRIS,
                             fillColor=HexColor("#FFFFFF"), fontSize=9, maxlen=6)
        c.setStrokeColor(TENUE); c.setLineWidth(0.5)
        c.line(m, y - alto_fila, xs[-1], y - alto_fila)
        y -= alto_fila
    c.showPage()
    c.save()
    return salida.getvalue()


def generar_excel(productos: list[Producto]) -> bytes:
    wb = Workbook()
    h = wb.active
    h.title = "Pedido"
    h.append(["cod_farmacia", "producto", "presentacion", "seccion", "minimo", "cantidad_a_pedir"])
    for p in productos:
        h.append([p.referencia, p.nombre, p.presentacion, p.seccion, _minimo(p), None])
    for celda in h[1]:
        celda.font = Font(bold=True, color="051923")
        celda.fill = PatternFill("solid", fgColor="7ED957" if celda.value == "cantidad_a_pedir" else "E5ECF1")
    for col, ancho in zip("ABCDEF", (16, 62, 14, 32, 14, 18)):
        h.column_dimensions[col].width = ancho
    for fila in range(2, len(productos) + 2):
        h[f"A{fila}"].number_format = "@"
    h.freeze_panes = "A2"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _cantidad(valor) -> int | None | str:
    """None si está vacío; 'error' si no es un entero mayor o igual a 0."""
    texto = str(valor if valor is not None else "").strip()
    if not texto:
        return None
    try:
        numero = float(texto.replace(",", "."))
    except ValueError:
        return "error"
    return int(numero) if numero >= 0 and numero == int(numero) else "error"


def leer(db: Session, contenido: bytes) -> tuple[dict[int, int], list[str]]:
    """Devuelve ({producto_id: cantidad}, errores) desde el PDF llenado o el Excel."""
    lineas, errores = {}, []
    if contenido.startswith(b"%PDF"):
        try:
            campos = PdfReader(io.BytesIO(contenido)).get_fields() or {}
        except Exception:
            return {}, ["No se pudo leer el PDF."]
        if not any(k.startswith(CAMPO) for k in campos):
            return {}, ["Este PDF no es el formato de Sinapsis (o se imprimió y escaneó). Use el formato descargado de la app."]
        validos = {p.id: p for p in productos_del_formato(db)}
        for nombre, campo in campos.items():
            if not nombre.startswith(CAMPO):
                continue
            pid = int(nombre[len(CAMPO):]) if nombre[len(CAMPO):].isdigit() else None
            q = _cantidad(campo.get("/V"))
            if q is None or q == 0:
                continue
            if pid not in validos:
                errores.append(f"Un producto del formato ya no existe o está inactivo (campo {nombre}).")
            elif q == "error":
                errores.append(f"{validos[pid].referencia}: la cantidad debe ser un número entero.")
            else:
                lineas[pid] = q
        return lineas, errores

    try:
        hoja = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True).worksheets[0]
    except Exception:
        return {}, ["El archivo debe ser el PDF o el Excel del formato de pedido."]
    filas = hoja.iter_rows(values_only=True)
    encabezado = [normalizar(str(c or "")).replace(" ", "_") for c in next(filas, [])]
    col_cod = next((encabezado.index(k) for k in ("cod_farmacia", "codigo_farmacia", "referencia") if k in encabezado), None)
    col_q = next((encabezado.index(k) for k in ("cantidad_a_pedir", "cantidad") if k in encabezado), None)
    if col_cod is None or col_q is None:
        return {}, ["El Excel debe tener las columnas cod_farmacia y cantidad_a_pedir."]
    por_ref = {p.referencia: p for p in productos_del_formato(db)}
    for n, fila in enumerate(filas, start=2):
        if not fila or col_q >= len(fila):
            continue
        q = _cantidad(fila[col_q])
        if q is None or q == 0:
            continue
        ref = normalizar_codigo(fila[col_cod])
        if ref not in por_ref:
            errores.append(f"Fila {n}: el cód. farmacia {ref or '(vacío)'} no existe o está inactivo.")
        elif q == "error":
            errores.append(f"Fila {n}: la cantidad debe ser un número entero.")
        else:
            lineas[por_ref[ref].id] = lineas.get(por_ref[ref].id, 0) + q
    return lineas, errores
