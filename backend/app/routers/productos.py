"""Catálogo: consulta, edición, carga inicial desde Excel, etiquetas y exportación."""
import io
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import HTMLResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, selectinload

from ..database import obtener_db
from ..models import Alias, Lote, Producto, Usuario
from ..schemas import ProductoEntrada, producto_json
from ..seguridad import solo_admin, usuario_actual
from ..services import catalogo, secciones
from ..services.etiquetas import hoja_etiquetas, qr_svg
from ..services.inventario import estado_vencimiento
from ..services.texto import normalizar, normalizar_codigo

router = APIRouter(prefix="/api/productos", tags=["Catálogo"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("", summary="Listar o buscar productos")
def listar(q: str = "", filtro: str = Query("todos", pattern="^(todos|bajo_minimo|qr)$"),
           limite: int = Query(50, le=500), desde: int = 0,
           _: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    consulta = select(Producto).where(Producto.activo.is_(True))
    if q:
        n, c = f"%{normalizar(q)}%", normalizar_codigo(q)
        consulta = consulta.where(or_(Producto.nombre_normalizado.like(n), Producto.referencia == c,
                                      Producto.seccion.ilike(f"%{q.strip()}%"),
                                      Producto.codigo_barras == c, Producto.codigo_qr == c))
    if filtro == "bajo_minimo":
        consulta = consulta.where(Producto.cantidad_actual < Producto.stock_minimo)
    elif filtro == "qr":
        consulta = consulta.where(Producto.metodo_identificacion == "qr")
    productos = db.scalars(consulta.order_by(Producto.nombre).offset(desde).limit(limite)).all()
    return [producto_json(p) for p in productos]


@router.get("/plantilla", summary="Descargar la plantilla de Excel para la carga inicial")
def plantilla(_: Usuario = Depends(usuario_actual)):
    return Response(catalogo.generar_plantilla(), media_type=XLSX,
                    headers={"Content-Disposition": 'attachment; filename="plantilla_catalogo_sinapsis.xlsx"'})


@router.post("/carga", summary="Validar o aplicar la carga del catálogo desde Excel (admin)")
async def carga(archivo: UploadFile = File(...), aplicar: bool = False,
                u: Usuario = Depends(solo_admin), db: Session = Depends(obtener_db)):
    contenido = await archivo.read()
    if len(contenido) > 15 * 1024 * 1024:
        raise HTTPException(413, "El archivo supera 15 MB.")
    filas, errores = catalogo.leer_excel(contenido)
    resultado = {"filas_validas": len(filas), "errores": errores, "aplicado": False}
    if aplicar and not errores and filas:
        aplicado = catalogo.aplicar_carga(db, filas, u.usuario)
        resultado.update(aplicado, aplicado=not aplicado["errores"])
    return resultado


@router.get("/etiquetas", response_class=HTMLResponse, summary="Hoja imprimible de etiquetas QR")
def etiquetas(ids: str = "", todos_qr: bool = True,
              _: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    consulta = select(Producto).where(Producto.activo.is_(True))
    if ids:
        lista = [int(i) for i in ids.split(",") if i.strip().isdigit()]
        consulta = consulta.where(Producto.id.in_(lista))
    elif todos_qr:
        consulta = consulta.where(Producto.metodo_identificacion == "qr")
    return hoja_etiquetas(db.scalars(consulta.order_by(Producto.nombre)).all())


@router.get("/exportar", summary="Exportar inventario y lotes a Excel")
def exportar(_: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Inventario"
    enc = ["Cód. farmacia", "Nombre", "Sección", "Presentación", "Código de barras", "Código QR",
           "Pruebas por unidad", "Cantidad actual", "Stock mínimo", "Se cuenta en", "Bajo mínimo"]
    hoja.append(enc)
    productos = db.scalars(select(Producto).options(selectinload(Producto.lotes))
                           .where(Producto.activo.is_(True)).order_by(Producto.nombre)).all()
    for p in productos:
        hoja.append([p.referencia, p.nombre, p.seccion, p.presentacion, p.codigo_barras or "", p.codigo_qr or "",
                     p.pruebas_por_unidad or "", p.cantidad_actual, p.stock_minimo,
                     "Pruebas" if p.pruebas_por_unidad else "Unidades",
                     "Sí" if p.cantidad_actual < p.stock_minimo else "No"])
    lotes = wb.create_sheet("Lotes y vencimientos")
    lotes.append(["Cód. farmacia", "Nombre", "Lote", "Cantidad ingresada (pruebas o unidades)", "Fecha ingreso",
                  "Fecha vencimiento", "Estado", "Origen"])
    etiqueta = {"vencido": "Vencido", "por_vencer": "Por vencer", "vigente": "Vigente", "sin_fecha": "Sin fecha"}
    for p in productos:
        for l in p.lotes:
            if l.estado == "activo":
                lotes.append([p.referencia, p.nombre, l.lote, l.cantidad, l.fecha_ingreso,
                              l.fecha_vencimiento, etiqueta[estado_vencimiento(l.fecha_vencimiento)], l.origen])
    for h in (hoja, lotes):
        for c in h[1]:
            c.font, c.fill = Font(bold=True, color="051923"), PatternFill("solid", fgColor="7ED957")
        for col in h.columns:
            h.column_dimensions[col[0].column_letter].width = 18
        h.freeze_panes = "A2"
    salida = io.BytesIO()
    wb.save(salida)
    nombre = f"sinapsis_inventario_{date.today().isoformat()}.xlsx"
    return Response(salida.getvalue(), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


@router.get("/{pid}", summary="Detalle de un producto con sus lotes activos")
def detalle(pid: int, _: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    p = db.get(Producto, pid)
    if not p:
        raise HTTPException(404, "Producto no encontrado.")
    return producto_json(p, detalle=True)


@router.get("/{pid}/qr.svg", summary="Código QR del producto")
def qr(pid: int, _: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    p = db.get(Producto, pid)
    if not p or not p.codigo_qr:
        raise HTTPException(404, "Producto no encontrado.")
    return Response(qr_svg(p.codigo_qr), media_type="image/svg+xml")


@router.post("", summary="Crear un producto (admin)")
def crear(datos: ProductoEntrada, _: Usuario = Depends(solo_admin), db: Session = Depends(obtener_db)):
    return producto_json(_guardar(db, Producto(cantidad_actual=0), datos))


@router.put("/{pid}", summary="Editar un producto (admin)")
def editar(pid: int, datos: ProductoEntrada, _: Usuario = Depends(solo_admin), db: Session = Depends(obtener_db)):
    p = db.get(Producto, pid)
    if not p:
        raise HTTPException(404, "Producto no encontrado.")
    return producto_json(_guardar(db, p, datos))


def _guardar(db: Session, p: Producto, d: ProductoEntrada) -> Producto:
    ref, codigo = normalizar_codigo(d.referencia), normalizar_codigo(d.codigo_barras) or None
    if db.scalar(select(Producto).where(Producto.referencia == ref, Producto.id != (p.id or 0))):
        raise HTTPException(400, "Ya existe otro producto con ese código de farmacia.")
    if codigo and db.scalar(select(Producto).where(Producto.codigo_barras == codigo, Producto.id != (p.id or 0))):
        raise HTTPException(400, "Ya existe otro producto con ese código de barras.")
    if p.id and p.cantidad_actual and d.pruebas_por_unidad != p.pruebas_por_unidad:
        raise HTTPException(400, "Para cambiar las pruebas por unidad, primero registre un conteo en 0 "
                                 "y luego vuelva a ingresar la existencia.")
    p.nombre, p.referencia, p.codigo_barras = d.nombre.strip(), ref, codigo
    p.pruebas_por_unidad, p.seccion = d.pruebas_por_unidad, secciones.asegurar(db, d.seccion)
    p.nombre_normalizado = normalizar(d.nombre)
    p.presentacion, p.stock_minimo = d.presentacion.strip(), d.stock_minimo
    p.metodo_identificacion = d.metodo_identificacion
    db.add(p)
    db.flush()
    if not p.codigo_qr:
        p.codigo_qr = f"SNP-{p.id:06d}"
    db.commit()
    return p


@router.delete("/{pid}/alias/{aid}", summary="Quitar un alias mal aprendido (admin)")
def quitar_alias(pid: int, aid: int, _: Usuario = Depends(solo_admin), db: Session = Depends(obtener_db)):
    a = db.get(Alias, aid)
    if not a or a.producto_id != pid:
        raise HTTPException(404, "Alias no encontrado.")
    db.delete(a)
    db.commit()
    return {"ok": True}


@router.post("/{pid}/lotes/{lid}/retirar", summary="Retirar un lote (vencido, dañado o agotado)")
def retirar_lote(pid: int, lid: int, motivo: str = Query("retirado", pattern="^(retirado|agotado)$"),
                 _: Usuario = Depends(usuario_actual), db: Session = Depends(obtener_db)):
    l = db.get(Lote, lid)
    if not l or l.producto_id != pid:
        raise HTTPException(404, "Lote no encontrado.")
    l.estado = motivo
    db.commit()
    return {"ok": True}
