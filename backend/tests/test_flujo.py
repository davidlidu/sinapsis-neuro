"""Prueba de punta a punta: carga inicial → identificación → ingreso → pedido → recepción → alertas."""
import io
import os
import tempfile
from datetime import date, timedelta

os.environ["SINAPSIS_DATA_DIR"] = tempfile.mkdtemp(prefix="sinapsis_test_")
os.environ["SINAPSIS_ADMIN_CLAVE"] = "clave-de-prueba"
os.environ.pop("OPENROUTER_API_KEY", None)

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from app.main import app
from app.services.catalogo import generar_plantilla
from app.services.texto import normalizar


@pytest.fixture(scope="module")
def c():
    with TestClient(app) as cliente:
        r = cliente.post("/api/auth/login", json={"usuario": "admin", "clave": "clave-de-prueba"})
        assert r.status_code == 200, r.text
        cliente.headers["Authorization"] = f"Bearer {r.json()['token']}"
        yield cliente


def _excel(filas):
    wb = load_workbook(io.BytesIO(generar_plantilla()))
    h = wb["Catalogo"]
    h.delete_rows(2, 2)
    for n, f in enumerate(filas, start=2):
        for col, valor in enumerate(f, start=1):
            h.cell(row=n, column=col, value=valor)
    s = io.BytesIO()
    wb.save(s)
    return s.getvalue()


def test_normalizar():
    assert normalizar("ACETAMINOFÉN 500MG Tabletas") == "acetaminofen 500 mg tab"


def test_sin_sesion_rechaza():
    with TestClient(app) as anon:
        assert anon.get("/api/productos").status_code == 401


def test_carga_valida_errores(c):
    malo = _excel([["", "R1", "", "", "", "", "x", 3, "", ""],
                   ["Uno", "R2", "", "", "", "", 1, 1, "", ""],
                   ["Dos", "R2", "", "", "", "", 1, 1, "31/02/2026", ""]])
    r = c.post("/api/productos/carga?aplicar=true", files={"archivo": ("c.xlsx", malo)}).json()
    assert not r["aplicado"]
    assert any("Fila 2" in e and "nombre vacío" in e for e in r["errores"])
    assert any("Fila 4" in e and "cód. farmacia repetido" in e for e in r["errores"])


def test_flujo_completo(c):
    vence = (date.today() + timedelta(days=10)).isoformat()
    archivo = _excel([
        ["Acetaminofén", "MED-0001", "7702001000011", "", "Paracetamol", "500 mg x 100 tab", 20, 35, vence, "L1"],
        ["Levetiracetam", "MED-0002", "", "nombre", "Keppra", "500 mg x 30 tab", 10, 4, "", ""],
        ["Jeringa desechable", "INS-0003", "", "qr", "", "5 ml", 50, 120, "", ""],
        ["Carbamazepina", "MED-0004", "", "", "", "200 mg x 30 tab", 15, 0, "", ""],
    ])
    r = c.post("/api/productos/carga?aplicar=true", files={"archivo": ("c.xlsx", archivo)}).json()
    assert r["aplicado"] and r["creados"] == 4, r

    # Identificación por código de barras y por QR propio
    p1 = c.get("/api/identificar/codigo/7702001000011").json()
    assert p1["referencia"] == "MED-0001"
    jeringa = c.get("/api/productos", params={"q": "jeringa"}).json()[0]
    assert c.get(f"/api/identificar/codigo/{jeringa['codigo_qr']}").json()["id"] == jeringa["id"]

    # Por texto con errores típicos de OCR y por alias
    cands = c.post("/api/identificar/texto", json={"texto": "LEVETIRACETAN 500MG"}).json()["candidatos"]
    assert cands[0]["referencia"] == "MED-0002"
    assert c.post("/api/identificar/texto", json={"texto": "keppra"}).json()["candidatos"][0]["referencia"] == "MED-0002"

    # Confirmar aprende la lectura como alias
    r = c.post("/api/identificar/confirmar", json={"producto_id": cands[0]["id"], "texto_leido": "LEVETlRACETAM 5OO"}).json()
    assert r["alias_aprendido"]

    # Ingreso manual con vencimiento
    lev = cands[0]["id"]
    r = c.post("/api/inventario/ingresos", json={"producto_id": lev, "cantidad": 2,
                                                  "fecha_vencimiento": "2020-01-01", "lote": "VIEJO"}).json()
    assert r["producto"]["cantidad_actual"] == 6

    alertas = c.get("/api/inventario/alertas").json()
    assert len(alertas["vencidos"]) == 1 and len(alertas["por_vencer"]) == 1
    assert {p["referencia"] for p in alertas["bajo_minimo"]} == {"MED-0002", "MED-0004"}

    # Pedido sugerido: bajo el mínimo, en cero o con lotes que vencen en los próximos 10 días (configurable)
    ped = c.post("/api/pedidos").json()
    sug = {i["producto"]["referencia"]: i["cantidad_sugerida"] for i in ped["items"]}
    motivos = {i["producto"]["referencia"]: i["motivo"] for i in ped["items"]}
    assert sug == {"MED-0001": 35, "MED-0002": 4, "MED-0004": 15}
    assert motivos["MED-0001"] == "Vence en 10 días" and motivos["MED-0004"] == "En cero"
    assert motivos["MED-0002"] == "Bajo el mínimo · Vencido"
    assert c.post("/api/pedidos").status_code == 400  # un solo borrador a la vez
    acet = next(i for i in ped["items"] if i["producto"]["referencia"] == "MED-0001")
    ped = c.put(f"/api/pedidos/{ped['id']}/items", json={"producto_id": acet["producto_id"], "cantidad": 0}).json()
    assert ped["total_items"] == 2  # se puede quitar lo que no se necesita
    ped = c.put(f"/api/pedidos/{ped['id']}/items", json={"producto_id": jeringa["id"], "cantidad": 30}).json()
    assert ped["total_items"] == 3

    # No se envía sin PDF válido
    assert c.post(f"/api/pedidos/{ped['id']}/enviar", files={"pdf": ("x.pdf", b"hola")}).status_code == 400
    ped = c.post(f"/api/pedidos/{ped['id']}/enviar", files={"pdf": ("pedido.pdf", b"%PDF-1.4 prueba")}).json()
    assert ped["estado"] == "enviado"
    assert c.get(f"/api/pedidos/{ped['id']}/pdf").content.startswith(b"%PDF")

    items = {i["producto"]["referencia"]: i for i in ped["items"]}
    # Llega completo uno, parcial otro, y la jeringa no llega
    c.post(f"/api/pedidos/{ped['id']}/recibir", json={"item_id": items["MED-0002"]["id"], "cantidad": 4,
                                                      "fecha_vencimiento": "2028-01-31", "lote": "K9"})
    r = c.post(f"/api/pedidos/{ped['id']}/recibir", json={"item_id": items["MED-0004"]["id"], "cantidad": 5}).json()
    assert r["item"]["estado"] == "parcial" and r["pedido"]["estado"] == "recibido_parcial"
    assert c.get(f"/api/productos/{lev}").json()["cantidad_actual"] == 10  # lo recibido se sumó

    ped = c.post(f"/api/pedidos/{ped['id']}/cerrar", json={}).json()
    assert ped["estado"] == "cerrado_con_pendientes"
    pend = c.get("/api/inventario/alertas").json()["pendientes"]
    assert {p["producto"]["referencia"] for p in pend} == {"MED-0004", "INS-0003"}
    assert all(not p["posible_desabastecimiento"] for p in pend)

    # Segundo pedido en que la carbamazepina tampoco llega → posible desabastecimiento
    ped2 = c.post("/api/pedidos").json()
    ped2 = c.post(f"/api/pedidos/{ped2['id']}/enviar", files={"pdf": ("p2.pdf", b"%PDF-1.4")}).json()
    c.post(f"/api/pedidos/{ped2['id']}/cerrar", json={})
    pend = c.get("/api/inventario/alertas").json()["pendientes"]
    assert any(p["producto"]["referencia"] == "MED-0004" and p["posible_desabastecimiento"] for p in pend)

    # Entrega tardía sobre un pedido cerrado con pendientes
    r = c.post(f"/api/pedidos/{ped['id']}/recibir", json={"item_id": items["INS-0003"]["id"], "cantidad": 30})
    assert r.status_code == 200 and r.json()["item"]["estado"] == "llego"

    # Conteo físico, historial, etiquetas y exportación
    assert c.post(f"/api/inventario/conteo/{lev}", json={"cantidad": 7}).json()["cantidad_actual"] == 7
    assert c.get(f"/api/inventario/movimientos/{lev}").json()[0]["tipo"] == "conteo"
    hoja = c.get("/api/productos/etiquetas").text
    assert "SNP-" in hoja and "Jeringa" in hoja
    assert c.get("/api/productos/exportar").content[:2] == b"PK"


def test_ocr_foto(c):
    img = Image.new("RGB", (900, 260), "white")
    d = ImageDraw.Draw(img)
    d.text((30, 40), "LEVETIRACETAM", fill="black", font=ImageFont.load_default(size=72))
    d.text((30, 150), "500 mg x 30 tabletas", fill="black", font=ImageFont.load_default(size=48))
    s = io.BytesIO()
    img.save(s, format="JPEG")
    r = c.post("/api/identificar/foto", files={"foto": ("f.jpg", s.getvalue(), "image/jpeg")})
    if r.status_code == 503:
        pytest.skip("Tesseract no instalado")
    datos = r.json()
    assert datos["candidatos"] and datos["candidatos"][0]["referencia"] == "MED-0002", datos


def test_ia_desactivada(c):
    assert c.get("/api/identificar/ia/estado").json()["habilitada"] is False
    r = c.post("/api/identificar/ia", files={"foto": ("f.jpg", b"x", "image/jpeg")})
    assert r.status_code == 503


def test_reactivos_por_pruebas(c):
    """Kit de 750 pruebas: ingresos en kits, existencias y mínimo en pruebas, pedido redondeado a kits."""
    from app.models import Producto
    from app.services.pedidos import cantidad_sugerida

    archivo = _excel([["ALB2 X 750 PRUEBAS", "LAB-ALB2", "", "", "", "Kit", 150, 0, "", "", "Química", 750, "si"],
                      ["CALIBRADOR F.A.S.", "LAB-CFAS", "", "", "", "Kit", 0, 0, "", "", "Química", "", "si"],
                      ["FERRITIN VIEJA", "LAB-FERR", "", "", "", "Kit", 0, 0, "", "", "Inmunología", "", "no"]])
    r = c.post("/api/productos/carga?aplicar=true", files={"archivo": ("lab.xlsx", archivo)}).json()
    assert r["aplicado"] and r["creados"] == 3, r
    alb = c.get("/api/productos?q=LAB-ALB2").json()[0]
    assert alb["pruebas_por_unidad"] == 750 and alb["seccion"] == "Química" and alb["bajo_minimo"]
    assert not c.get("/api/productos?q=LAB-FERR").json()  # inactivo: no aparece

    # Ingreso de 2 kits = 1.500 pruebas; el lote también queda en pruebas.
    r = c.post("/api/inventario/ingresos", json={"producto_id": alb["id"], "cantidad": 2}).json()
    assert r["producto"]["cantidad_actual"] == 1500 and r["lote"]["cantidad"] == 1500
    # Conteo en pruebas: quedan 120, bajo el umbral de 150 → se piden kits enteros.
    p = c.post(f"/api/inventario/conteo/{alb['id']}", json={"cantidad": 120}).json()
    assert p["bajo_minimo"]
    assert cantidad_sugerida(Producto(stock_minimo=150, cantidad_actual=120, pruebas_por_unidad=750)) == 1
    assert cantidad_sugerida(Producto(stock_minimo=10, cantidad_actual=4)) == 6
    # Con existencias, no se deja cambiar las pruebas por unidad sin un conteo en 0.
    datos = {"nombre": alb["nombre"], "referencia": alb["referencia"], "pruebas_por_unidad": 500}
    assert c.put(f"/api/productos/{alb['id']}", json=datos).status_code == 400


def test_recarga_fija_existencia_inicial(c):
    """El catálogo se carga en 0 y las cantidades llegan después en otro Excel."""
    c.post("/api/productos/carga?aplicar=true", files={"archivo": ("a.xlsx", _excel([["TIRAS DE ORINA", "LAB-TIRAS", "", "", "", "Unidad", 2, 0, "", ""]]))})
    r = c.post("/api/productos/carga?aplicar=true", files={"archivo": ("b.xlsx", _excel([["TIRAS DE ORINA", "LAB-TIRAS", "", "", "", "Unidad", 2, 5, "2027-01-31", "L1"]]))}).json()
    assert r["actualizados"] == 1
    p = c.get("/api/productos?q=LAB-TIRAS").json()[0]
    assert p["cantidad_actual"] == 5
    # Ya con movimientos, una nueva recarga no pisa la existencia.
    c.post("/api/productos/carga?aplicar=true", files={"archivo": ("c.xlsx", _excel([["TIRAS DE ORINA", "LAB-TIRAS", "", "", "", "Unidad", 2, 99, "", ""]]))})
    assert c.get("/api/productos?q=LAB-TIRAS").json()[0]["cantidad_actual"] == 5


def test_notificaciones(c):
    clave = c.get("/api/notificaciones/clave").json()["clave"]
    assert len(clave) == 87  # punto P-256 sin comprimir (65 bytes) en base64url
    sub = {"endpoint": "https://push.example.com/abc", "keys": {"p256dh": "x", "auth": "y"}}
    assert c.post("/api/notificaciones/suscribir", json=sub).json()["ok"]
    assert c.post("/api/notificaciones/desuscribir", json={"endpoint": sub["endpoint"]}).json()["ok"]
    assert c.post("/api/notificaciones/probar").status_code == 400  # ya no tiene dispositivos
    from app.services.notificaciones import texto_resumen
    assert texto_resumen({"vencidos": 1, "por_vencer": 0, "bajo_minimo": 3, "pendientes": 0}) == "1 lote vencido · 3 productos bajo el mínimo"
    assert texto_resumen({"vencidos": 0, "por_vencer": 0, "bajo_minimo": 0, "pendientes": 0}) is None


def test_configuracion_y_formato_de_pedido(c):
    """Días de vencimiento configurables; formato en blanco (PDF y Excel) que se llena y se sube."""
    from pypdf import PdfReader, PdfWriter

    assert c.get("/api/ajustes").json()["valores"]["dias_vencimiento_pedido"] == 10
    assert c.put("/api/ajustes", json={"dias_vencimiento_pedido": 500}).status_code == 422
    assert c.put("/api/ajustes", json={"dias_vencimiento_pedido": 15}).json()["valores"]["dias_vencimiento_pedido"] == 15
    c.put("/api/ajustes", json={"dias_vencimiento_pedido": 10})

    # Excel: se llena cantidad_a_pedir y se sube.
    wb = load_workbook(io.BytesIO(c.get("/api/pedidos/formato?tipo=xlsx").content))
    h = wb.active
    assert [x.value for x in h[1]][:1] == ["cod_farmacia"] and h.max_row > 2
    filas = {h.cell(row=n, column=1).value: n for n in range(2, h.max_row + 1)}
    h.cell(row=filas["LAB-ALB2"], column=6, value=3)
    h.cell(row=filas["LAB-TIRAS"], column=6, value="x")
    s = io.BytesIO(); wb.save(s)
    r = c.post("/api/pedidos/formato", files={"archivo": ("f.xlsx", s.getvalue())})
    assert r.status_code == 400 and "entero" in r.json()["detail"]
    h.cell(row=filas["LAB-TIRAS"], column=6).value = None
    s = io.BytesIO(); wb.save(s)
    r = c.post("/api/pedidos/formato?reemplazar=true", files={"archivo": ("f.xlsx", s.getvalue())})
    assert r.status_code == 200, r.text
    ped = r.json()
    assert [(i["producto"]["referencia"], i["cantidad_pedida"], i["motivo"]) for i in ped["items"]] == [("LAB-ALB2", 3, "Formato manual")]

    # PDF con campos: se llena en un lector de PDF y se sube; reemplaza el borrador si se confirma.
    pdf = c.get("/api/pedidos/formato?tipo=pdf").content
    assert pdf.startswith(b"%PDF")
    lector = PdfReader(io.BytesIO(pdf))
    campo = next(k for k in lector.get_fields() if k.startswith("pedir_"))
    w = PdfWriter(); w.append(lector)
    for pagina in w.pages:
        w.update_page_form_field_values(pagina, {campo: "2"}, auto_regenerate=False)
    s = io.BytesIO(); w.write(s)
    assert c.post("/api/pedidos/formato", files={"archivo": ("f.pdf", s.getvalue())}).status_code == 409
    ped2 = c.post("/api/pedidos/formato?reemplazar=true", files={"archivo": ("f.pdf", s.getvalue())}).json()
    assert len(ped2["items"]) == 1 and ped2["items"][0]["cantidad_pedida"] == 2
    borradores = [x for x in c.get("/api/pedidos").json() if x["estado"] == "borrador"]
    assert len(borradores) == 1 and borradores[0]["total_items"] == 1  # el borrador anterior se reemplazó
    c.delete(f"/api/pedidos/{ped2['id']}")


def test_secciones(c):
    """Lista de secciones: el producto crea la suya, renombrar arrastra los productos, eliminar los mueve."""
    r = c.post("/api/productos", json={"nombre": "Medio agar sangre", "referencia": "SEC-1", "stock_minimo": 0,
                                       "seccion": "  Microbiología  "})
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    lista = {x["nombre"]: x for x in c.get("/api/secciones").json()}
    assert lista["Microbiología"]["productos"] == 1
    assert c.post("/api/secciones", json={"nombre": "microbiología"}).status_code == 400  # repetida
    otra = c.post("/api/secciones", json={"nombre": "Hematología · Prueba"}).json()
    sid = lista["Microbiología"]["id"]
    assert c.put(f"/api/secciones/{sid}", json={"nombre": "Hematología · Prueba"}).status_code == 400
    c.put(f"/api/secciones/{sid}", json={"nombre": "Microbiología · Medios"})
    assert [p["id"] for p in c.get("/api/productos", params={"seccion": "Microbiología · Medios"}).json()] == [pid]
    assert c.get(f"/api/productos/{pid}").json()["seccion"] == "Microbiología · Medios"
    assert c.delete(f"/api/secciones/{sid}?mover_a=Hematología · Prueba").json()["ok"]
    assert c.get(f"/api/productos/{pid}").json()["seccion"] == "Hematología · Prueba"
    assert "Microbiología · Medios" not in {x["nombre"] for x in c.get("/api/secciones").json()}
    c.delete(f"/api/secciones/{otra['id']}")
    assert c.get(f"/api/productos/{pid}").json()["seccion"] == ""
